"""Settings tab save round-trips for active profile controls."""
from __future__ import annotations

import copy
import os
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from PyQt6.QtWidgets import QApplication

import config
from config import APP_SETTINGS


@pytest.fixture
def isolated_settings(tmp_path, monkeypatch):
    original = copy.deepcopy(APP_SETTINGS)
    APP_SETTINGS.clear()
    APP_SETTINGS.update(copy.deepcopy(config.DEFAULT_SETTINGS))
    monkeypatch.setattr(config, "SETTINGS_PATH", tmp_path / "settings.yaml")
    monkeypatch.setattr(config, "LEGACY_SETTINGS_PATH", tmp_path / "legacy.yaml")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    yield app, tmp_path / "settings.yaml"
    APP_SETTINGS.clear()
    APP_SETTINGS.update(original)


def test_tab_save_publishes_only_after_persistence(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    _app, target = isolated_settings
    tab = TabAnalysisSettings("flt3")
    signals = []
    tab.settings_saved.connect(signals.append)
    tab.default_input.setText("C:/new-input")
    with patch.object(config.os, "replace", side_effect=PermissionError("denied")):
        assert tab.save() is False
    assert APP_SETTINGS["analyses"]["flt3"]["batch"]["base_input_dir"] != "C:/new-input"
    assert tab.default_input.text() == "C:/new-input"
    assert signals == []
    assert "denied" in tab.status_lbl.text()
    assert not target.exists()
    assert tab.save() is True
    assert APP_SETTINGS["analyses"]["flt3"]["batch"]["base_input_dir"] == "C:/new-input"
    assert signals == ["flt3"]
    assert config.load_settings(target)["analyses"]["flt3"]["batch"]["base_input_dir"] == "C:/new-input"


def test_tab_save_respects_operation_guard(isolated_settings):
    from PyQt6.QtWidgets import QWidget

    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    _app, target = isolated_settings
    class GuardedWindow(QWidget):
        def settings_changes_allowed(self):
            return False

    parent = GuardedWindow()
    tab = TabAnalysisSettings("flt3", parent)
    tab.default_input.setText("C:/blocked")
    assert tab.save() is False
    assert not target.exists()
    assert APP_SETTINGS["general"]["author"] == "OUS"


def test_profile_saves_do_not_overwrite_global_or_other_profile(isolated_settings):
    from gui_qt.tabs.tab_app_settings import TabAppSettings
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    app_tab = TabAppSettings()
    app_tab.author.setText("Changed author")
    assert app_tab.save() is True
    clonality = TabAnalysisSettings("clonality")
    clonality.default_input.setText("C:/clonality")
    assert clonality.save() is True
    flt3 = TabAnalysisSettings("flt3")
    flt3.default_input.setText("C:/flt3")
    assert flt3.save() is True
    assert APP_SETTINGS["general"]["author"] == "Changed author"
    assert APP_SETTINGS["analyses"]["clonality"]["batch"]["base_input_dir"] == "C:/clonality"


def test_profile_validation_preserves_edits(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    tab = TabAnalysisSettings("clonality")
    tab.patient_regex.setText("[")
    assert tab.save() is False
    assert tab.patient_regex.text() == "["
    assert "invalid" in tab.status_lbl.text()
    tab.patient_regex.setText("patient-(\\d+)")
    tab.tracking_excel_path.setText("report.csv")
    assert tab.save() is False
    assert ".xlsx" in tab.status_lbl.text()


def test_app_qc_validation_and_engine_is_read_only(isolated_settings):
    from gui_qt.tabs.tab_app_settings import TabAppSettings

    tab = TabAppSettings()
    tab.d_min_r2_warn.setValue(0.999)
    tab.d_min_r2_ok.setValue(0.990)
    assert tab.save() is False
    assert "less than or equal" in tab.status_lbl.text()
    assert not hasattr(tab, "chk_use_rust_engine")
    assert "engine" in tab.engine_status.text().lower()


def test_engine_status_reports_available_and_missing(monkeypatch):
    from gui_qt.tabs.tab_app_settings import native_engine_status

    monkeypatch.setitem(
        sys.modules,
        "fraggler_native",
        SimpleNamespace(is_available=lambda: True, version="1.2.3"),
    )
    assert native_engine_status() == "In-process Rust engine available (version 1.2.3)."
    monkeypatch.setitem(
        sys.modules,
        "fraggler_native",
        SimpleNamespace(is_available=lambda: False),
    )
    assert native_engine_status() == (
        "In-process Rust engine unavailable. Check the installed engine or configured workflow."
    )


def test_dirty_profile_is_not_overwritten_by_refresh(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    tab = TabAnalysisSettings("flt3")
    tab.default_input.setText("C:/unsaved")
    assert tab.is_dirty()
    APP_SETTINGS["analyses"]["flt3"]["batch"]["base_input_dir"] = "C:/background"
    tab.refresh_from_settings()
    assert tab.default_input.text() == "C:/unsaved"
    assert "Unsaved" in tab.dirty_lbl.text()


def test_advanced_profile_fields_are_collapsible(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    tab = TabAnalysisSettings("clonality")
    assert tab.interpretation_card.title() == "Rule-based interpretation"
    assert tab.chk_clonality_interpretation.text() == "Enable rule-based interpretation"
    assert tab.interpretation_card.isCheckable()
    assert not tab.interpretation_card.isChecked()
    assert tab.advanced_content.isHidden()
    assert not hasattr(tab, "clonality_model_path")
    assert not hasattr(tab, "_ml_status_label")
    assert not hasattr(tab, "chk_clonality_learning")
    assert not hasattr(tab, "clonality_learning_output_dir")
    assert not hasattr(tab, "_browse_clonality_model_path")
    assert not hasattr(tab, "_refresh_ml_status")
    tab.interpretation_card.setChecked(True)
    assert not tab.advanced_content.isHidden()


def test_disabled_patient_grouping_ignores_unused_invalid_regex(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    tab = TabAnalysisSettings("flt3")
    tab.chk_agg_pat.setChecked(False)
    tab.patient_regex.setText("[")
    assert tab.save() is True


def test_rule_based_interpretation_round_trip(isolated_settings):
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    _app, target = isolated_settings
    profile = APP_SETTINGS["analyses"]["clonality"]
    profile["interpretation"]["enabled"] = False
    profile["interpretation"]["model_path"] = "C:/legacy-models"
    profile["learning"] = {"enabled": True, "output_dir": "C:/legacy-learning"}

    tab = TabAnalysisSettings("clonality")
    tab.chk_clonality_interpretation.setChecked(True)
    assert tab.save() is True

    reloaded = config.load_settings(target)
    saved_profile = reloaded["analyses"]["clonality"]
    assert saved_profile["interpretation"] == {"enabled": True}
    assert "model_path" not in saved_profile["interpretation"]
    assert "learning" not in saved_profile


@pytest.mark.parametrize(
    "analysis_id, field_name, browse_method",
    [
        ("clonality", "tracking_excel_path", "_browse_excel_path"),
        ("flt3", "tracking_excel_path", "_browse_excel_path"),
        ("general", "tracking_excel_path", "_browse_excel_path"),
        ("clonality", "global_tracking_excel_path", "_browse_global_tracking_excel_path"),
        ("flt3", "global_tracking_excel_path", "_browse_global_tracking_excel_path"),
    ],
)
def test_selecting_and_saving_existing_tracking_workbook_preserves_contents(
    isolated_settings, monkeypatch, analysis_id, field_name, browse_method
):
    from openpyxl import Workbook, load_workbook
    from PyQt6.QtWidgets import QDialog, QFileDialog

    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    _app, settings_path = isolated_settings
    workbook_path = settings_path.parent / "existing-tracking.xlsx"
    workbook = Workbook()
    workbook.active.append(["Patient", "Note", "Formula"])
    workbook.active.append(["existing-patient", "Keep this history", "=1+2"])
    workbook.save(workbook_path)
    original_bytes = workbook_path.read_bytes()

    def select_existing(dialog):
        assert dialog.acceptMode() == QFileDialog.AcceptMode.AcceptOpen
        dialog.selectFile(str(workbook_path))
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(QFileDialog, "exec", select_existing)
    tab = TabAnalysisSettings(analysis_id)
    getattr(tab, browse_method)()
    assert Path(getattr(tab, field_name).text()) == workbook_path
    assert workbook_path.read_bytes() == original_bytes
    assert tab.save() is True
    saved_batch = config.load_settings(settings_path)["analyses"][analysis_id]["batch"]
    assert Path(saved_batch[field_name]) == workbook_path
    assert workbook_path.read_bytes() == original_bytes
    assert load_workbook(workbook_path).active["C2"].value == "=1+2"


def test_tracking_workbook_selection_supports_new_path_and_cancel(isolated_settings, monkeypatch):
    from PyQt6.QtWidgets import QDialog, QFileDialog

    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    _app, settings_path = isolated_settings
    tab = TabAnalysisSettings("flt3")
    selected = settings_path.parent / "new-tracking"

    def select_new(dialog):
        dialog.selectFile(str(selected))
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(QFileDialog, "exec", select_new)
    tab._browse_excel_path()
    expected = selected.with_suffix(".xlsx")
    assert Path(tab.tracking_excel_path.text()) == expected
    assert not expected.exists()
    monkeypatch.setattr(QFileDialog, "exec", lambda _dialog: QDialog.DialogCode.Rejected)
    tab._browse_excel_path()
    assert Path(tab.tracking_excel_path.text()) == expected
    assert tab.save() is True
    assert not expected.exists()


@pytest.mark.parametrize("analysis_id", ["clonality", "flt3", "general", "app"])
def test_settings_panel_has_one_keyboard_save_action(isolated_settings, analysis_id):
    from PyQt6.QtCore import Qt
    from PyQt6.QtTest import QTest
    from PyQt6.QtWidgets import QPushButton

    from gui_qt.tabs.tab_app_settings import TabAppSettings
    from gui_qt.tabs.tab_settings import TabAnalysisSettings

    app, settings_path = isolated_settings
    tab = TabAppSettings() if analysis_id == "app" else TabAnalysisSettings(analysis_id)
    tab.resize(960, 900)
    tab.show()
    try:
        app.processEvents()
        save_actions = [
            button for button in tab.findChildren(QPushButton)
            if button.isVisible() and button.text().startswith("Save ")
        ]
        assert len(save_actions) == 1
        save_action = save_actions[0]
        field = tab.author if analysis_id == "app" else tab.default_output
        field.setText("Keyboard author" if analysis_id == "app" else str(settings_path.parent / "output"))
        field.setFocus()
        for _ in range(35):
            if app.focusWidget() is save_action:
                break
            QTest.keyClick(app.focusWidget(), Qt.Key.Key_Tab)
            app.processEvents()
        assert app.focusWidget() is save_action, "The save action must be reachable with Tab"
        signals = []
        tab.settings_saved.connect(signals.append)
        QTest.keyClick(save_action, Qt.Key.Key_Space)
        app.processEvents()
        assert signals == [analysis_id]
        assert "saved" in tab.status_lbl.text()
        assert not tab.is_dirty()
        persisted = config.load_settings(settings_path)
        if analysis_id == "app":
            assert persisted["general"]["author"] == "Keyboard author"
        else:
            assert persisted["analyses"][analysis_id]["batch"]["output_base"] == field.text()
    finally:
        tab.close()

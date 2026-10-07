from __future__ import annotations

import copy
from zipfile import BadZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from config import APP_SETTINGS
from core.analyses.clonality import tracking_excel
from core.analyses.flt3 import qc_tracker


def _entry(analysis: str, specimen: str, well: str) -> dict:
    assay = "FR1" if analysis == "clonality" else "FLT3-ITD"
    return {
        "fsa": None,
        "file_name": f"{specimen}_{assay}__260526_{well}_H9TEST01.fsa",
        "source_run_dir": "historic_run",
        "assay": assay,
        "dit": specimen if specimen.startswith("26OUM") else "",
        "specimen_id": specimen,
        "group": "negative_control" if specimen == "NTC" else "sample",
        "run_date": "2026-05-26",
        "well_id": well,
        "ladder": "ROX400HD" if analysis == "clonality" else "GS500ROX",
        "ladder_qc_status": "ok",
        "ladder_fit_strategy": "linear",
        "ladder_expected_step_count": 16,
        "ladder_fitted_step_count": 16,
        "ladder_r2": 0.9999,
        "primary_peak_channel": "DATA1",
        "peaks_by_channel": {
            "DATA1": pd.DataFrame(columns=["label", "basepairs", "peaks", "area", "keep"]),
        },
    }


def _update(analysis: str, path, entries) -> None:
    if analysis == "clonality":
        tracking_excel.update_clonality_tracking_workbook(path, entries)
    else:
        qc_tracker.update_flt3_npm1_qc_tracker_workbook(path, entries)


@pytest.mark.parametrize("analysis", ["clonality", "flt3"])
@pytest.mark.parametrize("layout", ["current", "legacy_run", "split", "split_without_keys"])
def test_selected_existing_workbook_keeps_history_and_operator_content(
    tmp_path, analysis, layout,
):
    path = tmp_path / "selected_tracking.xlsx"
    patient = _entry(analysis, "26OUM00001", "A01")
    control = _entry(analysis, "NK" if analysis == "clonality" else "NTC", "B01")
    _update(analysis, path, [patient, control])

    workbook = load_workbook(path)
    patient_sheet = workbook["Patient_Runs"]
    user_column = patient_sheet.max_column + 1
    patient_sheet.cell(1, user_column, "OperatorComment")
    patient_sheet.cell(2, user_column, "Keep this reviewed sample")
    patient_sheet.cell(1, user_column + 1, "OperatorFormula")
    patient_sheet.cell(2, user_column + 1, "=LEN(C2)")
    notes = workbook.create_sheet("OperatorNotes")
    notes["A1"] = "Original notes"
    notes["B1"] = "=1+2"
    if layout == "legacy_run":
        workbook["Runs"].title = "Run"
    elif layout in {"split", "split_without_keys"}:
        del workbook["Runs"]
        if layout == "split_without_keys":
            for name in ("Patient_Runs", "Control_Runs"):
                sheet = workbook[name]
                for cell in sheet[1]:
                    if cell.value == "IdentityKey":
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(row, cell.column).value = None
    workbook.save(path)
    workbook.close()

    refreshed = copy.deepcopy(patient)
    refreshed["ladder_qc_status"] = "review_required"
    added = _entry(analysis, "26OUM00002", "A02")
    _update(analysis, path, [refreshed, added])
    _update(analysis, path, [refreshed, added])

    runs = pd.read_excel(path, sheet_name="Runs")
    assert set(runs["File"]) == {patient["file_name"], control["file_name"], added["file_name"]}
    assert len(runs) == 3
    patients = pd.read_excel(path, sheet_name="Patient_Runs")
    assert len(patients) == 2
    original = patients.loc[patients["File"].eq(patient["file_name"])].iloc[0]
    assert original["LadderQC"] == "review_required"
    assert original["OperatorComment"] == "Keep this reviewed sample"
    controls = pd.read_excel(path, sheet_name="Control_Runs")
    assert list(controls["File"]) == [control["file_name"]]
    workbook = load_workbook(path, data_only=False)
    assert workbook["Patient_Runs"].cell(2, user_column + 1).value == "=LEN(C2)"
    assert workbook["OperatorNotes"]["A1"].value == "Original notes"
    assert workbook["OperatorNotes"]["B1"].value == "=1+2"
    if layout == "legacy_run":
        assert workbook["Run"].max_row == 3
    workbook.close()


@pytest.mark.parametrize("analysis", ["clonality", "flt3"])
def test_history_read_failure_keeps_selected_workbook_unchanged(tmp_path, monkeypatch, analysis):
    path = tmp_path / "selected_tracking.xlsx"
    original = _entry(analysis, "26OUM00001", "A01")
    _update(analysis, path, [original])
    before = path.read_bytes()

    read_excel = pd.read_excel
    failed = False

    def fail_read(*args, **kwargs):
        nonlocal failed
        if not failed and kwargs.get("sheet_name") == "Runs":
            failed = True
            raise PermissionError("Workbook cannot be read now")
        return read_excel(*args, **kwargs)

    monkeypatch.setattr(pd, "read_excel", fail_read)
    with pytest.raises(PermissionError, match="cannot be read"):
        _update(analysis, path, [_entry(analysis, "26OUM00002", "A02")])
    assert path.read_bytes() == before
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_clonality_existing_identity_survives_settings_or_machine_salt_change(tmp_path, monkeypatch):
    batch = APP_SETTINGS["analyses"]["clonality"]["batch"]
    monkeypatch.setitem(batch, "tracking_identity_salt", "original-machine")
    path = tmp_path / "selected_tracking.xlsx"
    patient = _entry("clonality", "26OUM00001", "A01")
    _update("clonality", path, [patient])
    old_key = pd.read_excel(path, sheet_name="Runs").iloc[0]["IdentityKey"]
    workbook = load_workbook(path)
    sheet = workbook["Patient_Runs"]
    column = sheet.max_column + 1
    sheet.cell(1, column, "OperatorComment")
    sheet.cell(2, column, "Preserve across PCs")
    workbook.save(path)
    workbook.close()

    monkeypatch.setitem(batch, "tracking_identity_salt", "different-machine")
    _update("clonality", path, [patient])

    runs = pd.read_excel(path, sheet_name="Runs")
    patients = pd.read_excel(path, sheet_name="Patient_Runs")
    assert len(runs) == 1
    assert runs.iloc[0]["IdentityKey"] == old_key
    assert patients.iloc[0]["OperatorComment"] == "Preserve across PCs"


@pytest.mark.parametrize("analysis", ["clonality", "flt3"])
def test_invalid_selected_workbook_is_never_replaced(tmp_path, analysis):
    path = tmp_path / "selected_tracking.xlsx"
    before = b"Workbook damaged during transfer"
    path.write_bytes(before)
    with pytest.raises(BadZipFile):
        _update(analysis, path, [_entry(analysis, "26OUM00001", "A01")])
    assert path.read_bytes() == before
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


@pytest.mark.parametrize("specimen, well, sheet_name", [
    ("26OUM00001", "A01", "Patient_Runs"),
    ("NTC", "B01", "Control_Runs"),
])
def test_flt3_repeated_operator_edits_and_clears_replace_previous_cached_notes(
    tmp_path, specimen, well, sheet_name,
):
    path = tmp_path / "selected_tracking.xlsx"
    entry = _entry("flt3", specimen, well)
    _update("flt3", path, [entry])
    for status, note in (
        ("Reviewed", "First review note"),
        ("Follow-up", "Corrected review note"),
        (None, None),
    ):
        workbook = load_workbook(path)
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["ReviewStatus"]).value = status
        sheet.cell(2, headers["TrackingNote"]).value = note
        workbook.save(path)
        workbook.close()
        _update("flt3", path, [entry])
        workbook = load_workbook(path)
        for name in ("Runs", sheet_name):
            sheet = workbook[name]
            headers = {cell.value: cell.column for cell in sheet[1]}
            assert sheet.cell(2, headers["ReviewStatus"]).value == status
            assert sheet.cell(2, headers["TrackingNote"]).value == note
        workbook.close()


@pytest.mark.parametrize("specimen, well, sheet_name", [
    ("26OUM00001", "A01", "Patient_Runs"),
    ("NTC", "B01", "Control_Runs"),
])
def test_flt3_built_in_operator_formulas_survive_repeated_updates(
    tmp_path, specimen, well, sheet_name,
):
    path = tmp_path / "selected_tracking.xlsx"
    entry = _entry("flt3", specimen, well)
    _update("flt3", path, [entry])
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["ReviewStatus"], '=IF(1=1,"Reviewed","Follow-up")')
    sheet.cell(2, headers["TrackingNote"], '=CONCAT("Checked ",C2)')
    workbook.save(path)
    workbook.close()
    _update("flt3", path, [entry])
    _update("flt3", path, [entry])
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[sheet_name]
    assert sheet.cell(2, headers["ReviewStatus"]).value == '=IF(1=1,"Reviewed","Follow-up")'
    assert sheet.cell(2, headers["TrackingNote"]).value == '=CONCAT("Checked ",C2)'
    workbook.close()

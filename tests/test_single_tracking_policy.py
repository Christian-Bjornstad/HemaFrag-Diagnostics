from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook


@pytest.mark.parametrize('analysis', ['clonality', 'flt3'])
def test_only_selected_workbook_is_resolved(analysis, tmp_path):
    from config import resolve_analysis_excel_output_path
    settings = {'analyses': {analysis: {'batch': {
        'tracking_excel_path': str(tmp_path / 'selected.xlsx'),
        'global_tracking_excel_path': str(tmp_path / 'old-master.xlsx'),
    }}}}
    assert resolve_analysis_excel_output_path(analysis, tmp_path, 'extra.xlsx', settings) == tmp_path / 'selected.xlsx'
    settings['analyses'][analysis]['batch']['tracking_excel_path'] = ''
    assert resolve_analysis_excel_output_path(analysis, tmp_path, 'extra.xlsx', settings) == tmp_path / 'old-master.xlsx'
    settings['analyses'][analysis]['batch']['global_tracking_excel_path'] = ''
    assert resolve_analysis_excel_output_path(analysis, tmp_path, 'extra.xlsx', settings) is None


@pytest.mark.parametrize('analysis', ['clonality', 'flt3'])
def test_qc_uses_selected_file_without_legacy_parallel_writer(monkeypatch, tmp_path, analysis):
    from core.runner import run_qc_job
    from core.run_context import RunContext
    from unittest.mock import Mock
    context = RunContext.create(analysis_id=analysis, settings={'analyses': {analysis: {
        'batch': {'tracking_excel_path': str(tmp_path / 'selected.xlsx')},
    }}})
    monkeypatch.setattr('core.pipeline.run_pipeline', lambda **kwargs: [{'file_name': 'PK.fsa'}])
    legacy = Mock()
    monkeypatch.setattr('core.qc.qc_excel.update_excel_trends', legacy)
    report = Mock()
    monkeypatch.setattr('core.qc.qc_html.build_qc_html', report)
    run_qc_job(tmp_path, tmp_path / 'output', 'qc.html', 'unwanted.xlsx', None, run_context=context)
    legacy.assert_not_called()
    assert report.call_args.args[3] == tmp_path / 'selected.xlsx'
    assert not list(tmp_path.rglob('*.xlsx'))


def test_tracking_uses_plain_filtered_ranges(tmp_path):
    from core.tracking_workbook_io import write_tracking_frames
    path = tmp_path / 'tracking.xlsx'
    frame = pd.DataFrame([{'IdentityKey': 'a', 'Assay': 'NPM1'}])
    write_tracking_frames(path, [('Runs', frame, ('IdentityKey',))])
    book = load_workbook(path)
    try:
        assert not book["Runs"].tables
        assert book["Runs"].auto_filter.ref == "A1:B2"
    finally:
        book.close()


def test_existing_table_schema_grows_without_invalid_column_metadata(tmp_path):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table
    from core.tracking_workbook_io import write_tracking_frames
    path = tmp_path / 'existing.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Runs'
    sheet.append(['IdentityKey', 'Assay'])
    sheet.append(['a', 'FR1'])
    sheet.add_table(Table(displayName='ExistingRuns', ref='A1:B2'))
    book.save(path)
    book.close()
    write_tracking_frames(path, [('Runs', pd.DataFrame([
        {'IdentityKey': 'a', 'Assay': 'FR1', 'Operator': 'Anne'},
    ]), ('IdentityKey',))])
    book = load_workbook(path)
    try:
        table = book['Runs'].tables['ExistingRuns']
        assert table.ref == 'A1:C2'
        assert [column.name for column in table.tableColumns] == ['IdentityKey', 'Assay', 'Operator']
    finally:
        book.close()


def test_operator_is_literal_text_and_captured_in_report(tmp_path):
    from core.html_reports._legacy import _create_html_header
    lines = []
    _create_html_header('26OUM00001', '2026', 1, tmp_path, lines,
                        display_name='Clonality', entries=[{'operator': 'Anne <Lab> & Co'}])
    html = '\n'.join(lines)
    assert 'Operatør: Anne &lt;Lab&gt; &amp; Co' in html
    assert 'Operatør: Anne <Lab>' not in html


def test_settings_migrates_legacy_master_and_allows_disabling(qapp, monkeypatch, tmp_path):
    import config
    from gui_qt.tabs.tab_settings import TabAnalysisSettings
    profile = config.APP_SETTINGS['analyses']['flt3']['batch']
    monkeypatch.setitem(profile, 'tracking_excel_path', '')
    monkeypatch.setitem(profile, 'global_tracking_excel_path', str(tmp_path / 'chosen.xlsx'))
    page = TabAnalysisSettings('flt3')
    assert page.tracking_excel_path.text() == str(tmp_path / 'chosen.xlsx')
    assert page.save()
    saved = config.load_settings(config.SETTINGS_PATH)['analyses']['flt3']['batch']
    assert saved['tracking_excel_path'] == str(tmp_path / 'chosen.xlsx')
    assert saved['global_tracking_excel_path'] == ''
    page.tracking_excel_path.clear()
    assert page.save()
    saved = config.load_settings(config.SETTINGS_PATH)
    assert config.resolve_analysis_excel_output_path('flt3', tmp_path, 'extra.xlsx', saved) is None
    assert not list(tmp_path.glob('*.xlsx'))


@pytest.mark.parametrize('analysis,assays', [
    ('flt3', ['FLT3-ITD', 'FLT3-D835', 'NPM1']),
    ('clonality', ['FR1', 'FR2', 'FR3', 'IGK', 'KDE', 'TCRbA', 'TCRbB', 'TCRbC',
                   'TCRgA', 'TCRgB', 'DHJH_D', 'DHJH_E', 'IKZF1', 'Ktr-albumin',
                   'IGHV Mix 1', 'IGHV Mix 2', 'SL']),
])
def test_all_assays_retain_history_and_operator_in_one_file(tmp_path, analysis, assays):
    if analysis == 'flt3':
        from core.analyses.flt3.qc_tracker import update_flt3_npm1_qc_tracker_workbook as update
    else:
        from core.analyses.clonality.tracking_excel import update_clonality_tracking_workbook as update
    path = tmp_path / 'selected.xlsx'
    entries = [{'file_name': f'26OUM00001_{assay}_A01.fsa', 'assay': assay,
                'source_run_dir': 'synthetic_run', 'dit': '26OUM00001',
                'operator': 'Test Operator', 'group': 'sample'} for assay in assays]
    if analysis == "flt3":
        from test_flt3_tracking_output import _entry
        entries = [dict(_entry(e["file_name"], assay=e["assay"]), operator="Test Operator") for e in entries]
    update(path, entries)
    update(path, entries)
    rows = pd.read_excel(path, sheet_name='Runs')
    assert set(rows['Assay']) == set(assays)
    assert len(rows) == len(assays)
    assert set(rows['Operator']) == {'Test Operator'}
    assert list(tmp_path.glob('*.xlsx')) == [path]
